"""One-off import of the previous team's spreadsheet (Inventario Tech Lab.xlsx) into places, items, assets and movements.

Usage: uv run python scripts/import_inventory.py <xlsx>            dry run: prints what it would write
       uv run python scripts/import_inventory.py <xlsx> --apply    writes it (service key); refuses a second run
The spreadsheet isn't trusted: every place starts with counted_at empty, so it shows as "never counted" in the office
until someone counts it. Student loan records (PlayStation sheet) are listed, not imported: check them first.
"""
import re
import sys
from collections import Counter, defaultdict

NOTE = "import: Inventario Tech Lab.xlsx"

# code, name, tier, parent. LAB, R15 and B2 exist already (migration 20260926100000).
PLACES = [
    ("LAB-VIT", "Vitrine (Tech Lab)", "fast", "LAB"),
    ("LAB-SRV", "Grey server cabinet (Tech Lab)", "fast", "LAB"),
    ("R15-L", "Room 15 left cabinet", "fast", "R15"),
    ("R15-R", "Room 15 right cabinet", "fast", "R15"),
    *[(f"R15-{s}-S{n}", f"Room 15 {side} cabinet, shelf {n}{' (top)' if n == 1 else ' (bottom)' if n == 5 else ''}", "fast", f"R15-{s}")
      for s, side in (("L", "left"), ("R", "right")) for n in range(1, 6)],
    ("R15-BOX-GREEN", "Green storage box (room 15)", "fast", "R15"),
    ("R15-BOX-SENS", "PascoCar sensor box (room 15)", "fast", "R15"),
    ("R15-BOX-CARS", "PascoCar cars box (room 15)", "fast", "R15"),
    *[(f"CUP{n}", f"Classroom cupboard {n}", "fast", None) for n in range(1, 9)],
]
STATIONARY_AT = {"LAB", "LAB-SRV"}           # fixed in the room: screens, servers, ceiling projectors, furniture
SKIP = {"computadores fixos"}                # the Computadores sheet lists these one by one
FURNITURE_SKIP = {"Arcades", "Ecras Lenovo", "Projectores de teto"}   # already counted in the Sala 13/14 sheet
TAG = re.compile(r"\b(GS|PF|LV|AS)\s*-?\s*(\d+)", re.I)
CONSOLE = re.compile(r"^\d\d\s*-\s*\d+\s*-\s*\d+$")
PAD = re.compile(r"^\d{6} [0-9A-Z] \d{7}$")


def place_code(text):
    """The place code for a spreadsheet location. Unknown spellings raise, so nothing is guessed."""
    t = re.sub(r"\s+", " ", str(text)).strip().lower()
    if t == "15":
        return "R15"
    if t in {"13", "14", "13/14", "13/14/15"}:
        return "LAB"
    if t.startswith("vitrine"):
        return "LAB-VIT"
    if "servidor cinza" in t:
        return "LAB-SRV"
    m = re.match(r"^15 ?- ?(\d) ?(esquerdo|direito)$", t)
    if m:
        return f"R15-{'L' if m[2] == 'esquerdo' else 'R'}-S{m[1]}"
    if "verde" in t:
        return "R15-BOX-GREEN"
    if "sensores" in t:
        return "R15-BOX-SENS"
    if "carros" in t:
        return "R15-BOX-CARS"
    raise ValueError(f"unknown location: {text!r}")


def qty(value):
    """A count from the Quantidade cell: '11 + 10' -> 21, '15 (7 Azuis + 8 Vermelhos)' -> 15, 'Varios (100+)' -> 100, 'Varias' -> 1."""
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value or "")
    outside = [int(n) for n in re.findall(r"\d+", re.sub(r"\([^)]*\)", "", text))]
    if outside:
        return sum(outside)
    inside = re.findall(r"\d+", text)
    return int(inside[0]) if inside else 1


def tags(identifier):
    """Legacy asset tags in an Identificador cell: 'GS 089+ 090 + 092' -> GS-089, GS-090, GS-092; 'PF 136 +... + PF 145' is a range."""
    text = str(identifier or "")
    found = TAG.findall(text)
    if not found:
        return []
    prefix = found[0][0].upper()
    if "..." in text and len(found) == 2:
        a, b = int(found[0][1]), int(found[1][1])
        return [f"{prefix}-{n:03d}" for n in range(a, b + 1)]
    out = []
    for part in re.split(r"[+,/]", re.sub(r"\([^)]*\)", "", text)):
        m = TAG.search(part) or re.match(r"^\s*(\d+)\s*$", part)
        if m:
            p, n = (m[1].upper(), m[2]) if m.re is TAG else (prefix, m[1])
            prefix = p
            out.append(f"{p}-{int(n):03d}" if p in {"GS", "PF"} else f"{p}-{int(n):02d}")
    return out


def blank(v):
    return v is None or str(v).strip() in {"", "-", "- ", ".-"}


class Inventory:
    """What the spreadsheet says: lines (item at place, qty), assets (one row each, by tag), and what was left out."""

    def __init__(self):
        self.lines = []      # dict(item, place, qty, note)
        self.assets = {}     # tag or serial -> dict(item, place, tag, serial, condition, note)
        self.notes = defaultdict(list)
        self.skipped = []

    def line(self, item, place, n, note=None):
        item = re.sub(r"\s+", " ", item).strip()
        self.lines.append({"item": item, "place": place, "qty": n})
        if note:
            self.notes[item].append(note)

    def asset(self, item, place, tag=None, serial=None, condition="ok", note=None):
        key = tag or serial
        a = self.assets.setdefault(key, {"item": item, "place": place, "tag": tag, "serial": None, "condition": condition, "note": None})
        a["serial"] = a["serial"] or serial
        a["note"] = "; ".join(x for x in (a["note"], note) if x) or None

    def kinds(self):
        return {l["item"]: "stationary" if l["place"] in STATIONARY_AT else "portable" for l in self.lines}

    def receives(self):
        """(item, place, qty) left after the assets, which are received one by one."""
        have = Counter((a["item"], a["place"]) for a in self.assets.values())
        want = Counter()
        for l in self.lines:
            want[(l["item"], l["place"])] += l["qty"]
        return [(i, p, n - have[(i, p)]) for (i, p), n in want.items() if n - have[(i, p)] > 0]


def read_rows(inv, ws, item_col, holder_col=None):
    for r in ws.iter_rows(values_only=True):
        name, n, where, obs, ident = r[item_col:item_col + 5]
        if blank(name) or name == "Item" or blank(where) or str(where).startswith(("Prateleiras", "1 = Topo")):
            continue
        if str(name).strip().lower() in SKIP:
            continue
        place = place_code(where)
        name = str(name).strip()
        found = [] if holder_col else tags(ident)
        extra = [] if isinstance(n, (int, float)) or n is None else [f"quantity in the sheet: {n}"]
        if isinstance(n, (int, float)) and len(found) > n:   # device and box tags mixed: tag them at the stocktake
            extra.append(f"tags in the sheet: {', '.join(found)}")
            found = []
        count = len(found) if found and not isinstance(n, (int, float)) else qty(n)
        if not blank(obs):
            extra.append(str(obs).strip())
        if holder_col:
            extra.append(f"belongs to {str(ident).strip()}")
        elif not blank(ident) and not found:
            extra.append(f"identifier: {str(ident).strip()}")
        inv.line(name, place, count, "; ".join(extra) or None)
        for t in found:
            inv.asset(name, place, tag=t)


def read_computers(inv, ws):
    for r in ws.iter_rows(values_only=True):
        pid = r[1]
        if blank(pid) or pid == "Computador ID" or not TAG.search(str(pid)):
            continue
        tag = tags(pid)[0]
        specs = "; ".join(str(c).strip() for c in (r[2], r[3], r[4], r[5]) if not blank(c))
        broken = "avaria" in str(r[7] or "").lower()
        inv.asset("Computador Tech Lab", "LAB", tag=tag, condition="broken" if broken else "ok",
                  note="; ".join(x for x in (specs, str(r[7] or "").strip()) if x) or None)
    inv.line("Computador Tech Lab", "LAB", sum(1 for a in inv.assets.values() if a["item"] == "Computador Tech Lab"))


def read_furniture(inv, ws):
    rows = {str(r[1]).strip(): r for r in ws.iter_rows(values_only=True) if r[1]}
    for name, n, size in zip(rows["Item"][2:], rows["Quantidade"][2:], rows["Dimensões"][2:]):
        if name and str(name).strip() not in FURNITURE_SKIP:
            inv.line(str(name).strip(), "LAB", qty(n), re.sub(r"\s+", " ", str(size)).strip() if size else None)


def read_playstations(inv, ws):
    """Consoles (column C) and controllers (columns F and G) by serial, placed by the block header above them."""
    console_block, pad_place = None, None
    for r in ws.iter_rows(values_only=True):
        head, pads = str(r[1] or ""), str(r[5] or "")
        if head.startswith(("Consolas", "Playstations de Alunos")):
            console_block = head
        if pads.startswith("Comandos PS4"):
            pad_place = "LAB-VIT" if "Vitrine" in pads else "R15-L-S3"
        if console_block and "Alunos" in console_block and r[2] is not None and r[1] != "Nome Aluno":
            inv.skipped.append(f"student loan: {r[1]} ({r[2]}): console {r[3]}, box {r[4]}, controller {r[5]}")
            continue
        serial = str(r[2] or "").strip()
        if CONSOLE.match(serial):
            if "Fora do IADE" in console_block:
                inv.skipped.append(f"console outside IADE: {serial} {r[3]}")
            else:
                item = "Consola PS4 Classic" if "Old" in console_block else "Consola PS4 Slim"
                found = tags(r[3])
                inv.asset(item, "LAB-VIT", tag=found[0] if found else None, serial=re.sub(r"\s+", "", serial),
                          note=f"{r[1]}; {str(r[3]).strip()}")
        for cell in (r[5], r[6]):
            if pad_place and PAD.match(str(cell or "").strip()):
                item = "Comandos PS4" if pad_place == "LAB-VIT" else "Commando PS4"
                inv.asset(item, pad_place, serial=str(cell).strip())


def read_workbook(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    inv = Inventory()
    read_rows(inv, wb["Inventario Sala 15"], 2)
    read_rows(inv, wb["Sala 1314"], 2)
    read_rows(inv, wb["Material de Docente"], 1, holder_col=True)
    read_computers(inv, wb["Computadores Tech Lab"])
    read_furniture(inv, wb["Inventario Mobiliario"])
    read_playstations(inv, wb["Playstations Records"])
    return inv


def report(inv):
    by_place = defaultdict(list)
    for i, p, n in inv.receives():
        by_place[p].append(f"{n} × {i}")
    for a in inv.assets.values():
        by_place[a["place"]].append(f"asset {a['tag'] or '(TL-auto)'} {a['serial'] or ''} {a['item']} {a['condition']}".rstrip())
    for p in sorted(by_place):
        print(f"\n[{p}]")
        for x in by_place[p]:
            print("  " + x)
    print(f"\n{len(inv.kinds())} items, {len(inv.assets)} assets, {len(inv.receives())} counted lines")
    if inv.skipped:
        print("\nNot imported (decide by hand):")
        for s in inv.skipped:
            print("  " + s)


def apply(inv):
    from db import connect, request, select
    db = connect()
    if select(db, "movements", {"select": "id", "note": f"eq.{NOTE}", "order": "id", "limit": "1"}):
        sys.exit("Already imported: movements with the import note exist.")
    back = {"Prefer": "return=representation"}
    places = {p["code"]: p["id"] for p in select(db, "places", {"select": "id,code", "order": "id"}) if p["code"]}
    for code, name, tier, parent in PLACES:
        if code not in places:
            row = {"code": code, "name": name, "kind": "storage", "tier": tier, "parent_id": places.get(parent)}
            places[code] = request(db, "POST", "places", {"select": "id"}, row, back)[0]["id"]
    notes = {i: "; ".join(dict.fromkeys(n)) for i, n in inv.notes.items()}
    rows = [{"name": i, "kind": k, "note": notes.get(i)} for i, k in inv.kinds().items()]
    request(db, "POST", "items", {"on_conflict": "name"}, rows, {"Prefer": "resolution=ignore-duplicates"})
    items = {i["name"]: i["id"] for i in select(db, "items", {"select": "id,name", "order": "id"})}
    moves = []
    for a in inv.assets.values():
        row = {"item_id": items[a["item"]], "serial": a["serial"], "condition": a["condition"], "note": a["note"]}
        if a["tag"]:
            row["tag"] = a["tag"]
        aid = request(db, "POST", "assets", {"select": "id"}, row, back)[0]["id"]
        moves.append({"item_id": items[a["item"]], "asset_id": aid, "qty": 1, "to_place": places[a["place"]], "kind": "receive", "note": NOTE})
    for i, p, n in inv.receives():
        moves.append({"item_id": items[i], "asset_id": None, "qty": n, "to_place": places[p], "kind": "receive", "note": NOTE})
    request(db, "POST", "movements", None, moves)
    print(f"wrote {len(moves)} movements")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    inventory = read_workbook(sys.argv[1])
    report(inventory)
    if "--apply" in sys.argv:
        apply(inventory)
